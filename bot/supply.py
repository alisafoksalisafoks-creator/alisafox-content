"""Агент поставок Ozon FBO: куда и сколько отгрузить со своего склада.
Только чтение API, в кабинетах ничего не меняет. Остатки и скорость по кластерам берёт из restock.py,
сверху добавляет срок доставки до склада Ozon и кратность коробов.
Заказ поставщику здесь не считаем: решение о заказе опирается на чистую прибыль, это отдельный бот.

    python -m bot.supply [--out план.xlsx]

Настройки — bot/supply_settings.json, кратность коробов по артикулам — bot/supply_items.csv."""
import argparse
import csv
import json
import math
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from bot import restock
from bot.sources.moysklad import Catalog
from bot.sources.ozon import OzonClient

HERE = Path(__file__).parent
SETTINGS = HERE / "supply_settings.json"
ITEMS = HERE / "supply_items.csv"
DEFAULTS = {
    "delivery_days": 7,        # от отгрузки со своего склада до приёмки на Ozon
    "target_days": restock.TARGET_DAYS,  # запас на складе Ozon после приёмки
}


def settings():
    s = dict(DEFAULTS)
    if SETTINGS.exists():
        s.update({k: v for k, v in json.loads(SETTINGS.read_text("utf-8")).items() if not k.startswith("_")})
    return s


def items():
    """{корневой артикул: {box}} из supply_items.csv (article;box_qty). Пустое поле — кратность не задана."""
    out = {}
    if not ITEMS.exists():
        return out
    with ITEMS.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f, delimiter=";"):
            art = Catalog.root(row.get("article"))
            if not art:
                continue
            box = (row.get("box_qty") or "").strip()
            out[art] = {"box": int(float(box)) if box else None}
    return out


def ceil_to(qty, box):
    return math.ceil(qty / box) * box if box and qty > 0 else qty


def ship_plan(rows, catalog, cfg, boxes):
    """К строкам restock.plan добавляет: need — до target_days на момент приёмки, округлено вверх до короба;
    ship — сколько везти из своего остатка (сначала туда, где запас кончится раньше); order — не хватило своего."""
    lead = cfg["delivery_days"]
    groups = defaultdict(list)
    for r in rows:
        r["ship"] = r["order"] = 0
        r["box"] = (boxes.get(Catalog.root(r["offer_id"])) or {}).get("box")
        if r["status"] in ("urgent", "low") and r["ads"]:
            raw = r["ads"] * (lead + cfg["target_days"]) - r["stock"] - r["transit"]
            r["need"] = ceil_to(math.ceil(raw), r["box"]) if raw >= restock.MIN_NEED else 0
            if r["need"]:
                groups[Catalog.root(r["offer_id"])].append(r)
        else:
            r["need"] = 0
    for root, rs in groups.items():
        left = int(max(0, (catalog.get(root) or {}).get("stock", 0)))
        for r in sorted(rs, key=lambda r: (r["days"] or 0, -r["ads"])):
            box = r["box"]
            take = min(r["need"], left)
            if box and take < r["need"]:
                take = take // box * box  # неполный короб не отгружаем
            r["ship"], r["order"] = take, r["need"] - take
            left -= take
    return rows


def build(oz=None, catalog=None):
    oz = oz or OzonClient()
    catalog = catalog or Catalog()
    cfg, boxes = settings(), items()
    day = date.today() - timedelta(days=1)
    skus = [r["sku"][0] for r in oz.analytics(day - timedelta(days=29), day, ["sku"], ("ordered_units",))
            if r["ordered_units"] > 0]
    rows = ship_plan(restock.plan(restock.cluster_stocks(oz, skus)), catalog, cfg, boxes)
    return {"cfg": cfg, "rows": rows, "clusters": restock.by_cluster(rows), "skus": len(skus), "boxes": len(boxes),
            "anomalies": catalog.anomalies()}


def to_excel(plan, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cfg, rows = plan["cfg"], plan["rows"]
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    red, orange, blue = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FCE4D6", "DDEBF7"))

    def sheet(ws, title, cols, data, widths=None):
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        return ws

    d = lambda x: round(x) if x is not None else "нет продаж"
    ship = [r for r in rows if r["need"]]
    ship.sort(key=lambda r: (r["cluster"], r["days"] or 0))
    ship_total = sum(r["ship"] for r in ship)

    ws = wb.active
    ws.title = "Сводка"
    lines = [
        ("План поставок Ozon FBO", ""),
        ("Дата расчёта", date.today().strftime("%d.%m.%Y")),
        ("Источник", "Ozon Seller API /v1/analytics/stocks (скорость — средняя за 28 дн), МойСклад /report/stock/all"),
        ("SKU с заказами за 30 дней", plan["skus"]),
        ("Складов, куда пора везти", len(plan["clusters"])),
        ("Отгрузить со своего склада, шт", ship_total),
        ("Не хватает своего на отгрузку, шт", sum(r["order"] for r in ship)),
        ("Заказ поставщику", "не считаем: решение по чистой прибыли — отдельный бот"),
        ("", ""),
        ("Правила", ""),
        ("Везём, если запас на складе меньше, дн", restock.MIN_DAYS),
        ("Везём до запаса на момент приёмки, дн", cfg["target_days"]),
        ("Перебор, если запас больше, дн", restock.EXCESS_DAYS),
        ("Меньше этого на склад не везём, шт", restock.MIN_NEED),
        ("Доставка до приёмки Ozon, дн", f"{cfg['delivery_days']} (допущение, уточнить)"),
        ("Кратность коробов задана для артикулов", f"{plan['boxes']} (пока без кратности: шт не округлены)"
                                                    if not plan["boxes"] else plan["boxes"]),
    ]
    if plan["anomalies"]:
        lines.append(("Аномалии МоегоСклада (минус в остатке/цене)", ", ".join(sorted(plan["anomalies"]))))
    for row in lines:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.cell(lines.index(("Правила", "")) + 1, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 70

    sheet(wb.create_sheet(), "По складам",
          ["Кластер", "SKU в поставке", "Нужно, шт", "Везти со своего, шт", "Не хватает, шт", "SKU с нулём"],
          [[c["cluster"], c["skus"], c["need"], c["ship"], c["need"] - c["ship"], c["urgent"]]
           for c in plan["clusters"]], [32, 14, 12, 18, 14, 12])

    ws = sheet(wb.create_sheet(), "Отгрузка",
               ["Кластер", "Артикул", "SKU", "Название", "Статус", "Остаток, шт", "В пути, шт", "Продажи/день",
                "Запас, дн", "Нужно, шт", "Короб, шт", "Везти, шт", "Не хватает, шт"],
               [[r["cluster"], r["offer_id"], int(r["sku"]), r["name"],
                 "ноль" if r["status"] == "urgent" else "мало", r["stock"], r["transit"], round(r["ads"], 2),
                 d(r["days"]), r["need"], r["box"] or "", r["ship"], r["order"]] for r in ship],
               [26, 14, 12, 50, 8, 10, 10, 10, 9, 10, 9, 10, 12])
    for row in ws.iter_rows(min_row=2):
        row[4].fill = red if row[4].value == "ноль" else orange

    ex = sorted((r for r in rows if r["status"] == "excess"), key=lambda r: -r["excess"])
    ws = sheet(wb.create_sheet(), "Перебор",
               ["Кластер", "Артикул", "SKU", "Название", "Остаток+путь, шт", "Продажи/день", "Запас, дн",
                "Лишнее сверх 60 дн, шт"],
               [[r["cluster"], r["offer_id"], int(r["sku"]), r["name"], r["stock"] + r["transit"],
                 round(r["ads"], 2), d(r["days"]), r["excess"]] for r in ex],
               [26, 14, 12, 50, 12, 10, 9, 14])
    for row in ws.iter_rows(min_row=2):
        row[7].fill = blue
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=f"ozon_supply_plan_{date.today()}.xlsx")
    path = ap.parse_args().out
    plan = build()
    to_excel(plan, path)
    print(f"Готово: {path}. Складов к поставке {len(plan['clusters'])}, "
          f"везти {sum(r['ship'] for r in plan['rows'])} шт, "
          f"не хватает своего {sum(r['order'] for r in plan['rows'])} шт")


if __name__ == "__main__":
    main()
