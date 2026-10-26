"""Позиции артикулов в поиске и каталоге Ozon: средняя позиция по дням (метрика position_category).

Позиция приходит тем же запросом /v1/analytics/data, что и продажи (store.daily), отдельных
вызовов API нет. Чем меньше число, тем выше товар в выдаче; 0 — в этот день товар не показывался.
"""
import time
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bot.sources.ozon import OzonClient

SPAN = 28  # дашборд артикула — за последние 28 дней, как график в карточке
FLAT = 5  # изменение меньше 5% — «≈», это обычный разброс средней позиции
SUPER_TTL = 3600


_super = {"at": 0.0, "skus": None}


def super_skus(oz: OzonClient):
    """SKU с признаком Ozon «Супер-товар» (is_super) — у нас это новинки. Кэш на час."""
    if _super["skus"] is not None and time.time() - _super["at"] < SUPER_TTL:
        return _super["skus"]
    ids, last = [], ""
    while True:
        res = oz.post("/v3/product/list", {"filter": {"visibility": "ALL"}, "limit": 1000, "last_id": last})["result"]
        ids += [str(i["product_id"]) for i in res["items"]]
        last = res.get("last_id")
        if len(res["items"]) < 1000 or not last:
            break
    out = []
    for i in range(0, len(ids), 1000):
        for x in oz.post("/v3/product/info/list", {"product_id": ids[i:i + 1000]}).get("items", []):
            if x.get("is_super") and x.get("sku") and not x.get("is_archived"):
                out.append({"sku": str(x["sku"]), "offer_id": x["offer_id"], "name": x["name"],
                            "type_id": str(x.get("type_id") or ""), "created": x.get("created_at", "")[:10]})
    out.sort(key=lambda x: x["created"], reverse=True)
    _super.update(at=time.time(), skus=out)
    return out


def series(db, sku, start: date, end: date):
    """[(день, позиция или None, показы)] за каждый день периода."""
    got = {r[0]: (r[1] or None, r[2]) for r in db.execute(
        "select day, position_category, hits_view from daily where sku = ? and day between ? and ?",
        (str(sku), str(start), str(end)))}
    n = (end - start).days + 1
    return [(start + timedelta(days=i), *got.get(str(start + timedelta(days=i)), (None, 0))) for i in range(n)]


def change(cur, old):
    """Сколько мест прибавил товар: плюс — поднялся (число позиции уменьшилось)."""
    return None if cur is None or old is None else old - cur


def move(cur, old):
    """(подпись, тип): тип up | down | flat | none. Порог «≈» — 5% от прошлой позиции."""
    d = change(cur, old)
    if d is None:
        return "нет данных", "none"
    if abs(d) < max(1, old * FLAT / 100):
        return f"≈ {abs(d):.0f}", "flat"
    return (f"▲ {d:.0f}", "up") if d > 0 else (f"▼ {-d:.0f}", "down")


def mean(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def summary(db, items, day: date):
    """Строки списка: позиция за день, к предыдущему дню и к тому же дню неделю назад."""
    out = []
    for it in items:
        s = series(db, it["sku"], day - timedelta(days=7), day)
        cur, prev, week = s[-1][1], s[-2][1], s[0][1]
        out.append({**it, "pos": cur, "prev": prev, "week": week, "views": s[-1][2],
                    "d1": change(cur, prev), "d7": change(cur, week)})
    # сначала те, кто сильнее всего упал, затем выросшие; без данных — в конец
    return sorted(out, key=lambda r: (r["pos"] is None, r["d1"] if r["d1"] is not None else 0))


def sku_card(db, sku, day: date):
    s = series(db, sku, day - timedelta(days=SPAN - 1), day)
    prev = series(db, sku, day - timedelta(days=2 * SPAN - 1), day - timedelta(days=SPAN))
    pos = [p for _, p, _ in s]
    known = [(d, p) for d, p, _ in s if p is not None]
    best = min(known, key=lambda x: x[1]) if known else None
    worst = max(known, key=lambda x: x[1]) if known else None
    return {"series": s, "cur": pos[-1], "prev": pos[-2], "week": pos[-8], "avg": mean(pos),
            "avg_prev": mean(p for _, p, _ in prev), "best": best, "worst": worst,
            "views": sum(v or 0 for *_, v in s)}


def excel(db, items, start: date, end: date, path):
    """Excel за период: строка — артикул, столбцы — дни, в ячейке средняя позиция за день."""
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Позиции по дням"
    ws.append(["Артикул"] + [f"{d:%d.%m}" for d in days])
    for it in items:
        s = series(db, it["sku"], start, end)
        ws.append([it.get("offer_id") or it["sku"]] + [round(p) if p is not None else None for _, p, _ in s])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="F3F2EF")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = max([10] + [len(str(r[0].value or "")) + 2 for r in ws.iter_rows(max_col=1)])
    for i in range(2, len(days) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 7
    ws.append([])
    ws.append(["Средняя позиция в поиске и каталоге Ozon за день: чем меньше, тем выше товар. "
               "Пусто — в этот день товар не показывался. Источник: Ozon Seller API."])
    wb.save(path)
    return path
