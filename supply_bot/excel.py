"""Excel с распределением. Цифры те же, что бот показал на экране."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE7F0")
TOTAL = PatternFill("solid", fgColor="F2F2F2")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _boxes(ln):
    return ", ".join(f"{k}×{s} шт ({name})" for name, s, k in ln.boxes) or "—"


def _distribution(wb, plan, groups):
    """Лист распределения: {направление: [строки]} — по коробам, итог по каждому направлению."""
    ws = _sheet(wb, "Лист распределения",
                ["Направление", "Кластер", "Склад приёмки", "Артикул", "Кол-во, шт", "Короб", "Шт в коробе", "Коробов",
                 "Дата отгрузки", "Приоритет", "Пометки"],
                [], [30, 26, 28, 28, 11, 14, 11, 9, 13, 11, 50])
    for wh, lines in groups.items():
        for ln in sorted(lines, key=lambda x: (-["обычно", "наверх", "срочно"].index(x.priority), -x.qty)):
            if ln.boxes:
                for name, s, k in ln.boxes:
                    ws.append([wh, ln.cluster, ln.ship_to, ln.offer, s * k, name, s, k, plan.ship_date, ln.priority, "; ".join(ln.flags)])
            else:
                ws.append([wh, ln.cluster, ln.ship_to, ln.offer, ln.qty, "нет кванта", None, None, plan.ship_date, ln.priority,
                           "; ".join(ln.flags)])
        row = [f"Итого {wh}", "", "", f"{len(lines)} арт.", sum(x.qty for x in lines), "", "",
               sum(x.box_count for x in lines), "", "", ""]
        ws.append(row)
        for c in ws[ws.max_row]:
            c.font, c.fill = BOLD, TOTAL
    for c in ws["I"][1:]:
        c.number_format = "DD.MM.YYYY"


def _need(wb, lines, title="Потребность по направлениям"):
    """Весь расчёт по строкам: из чего получилась потребность и сколько везём."""
    _sheet(wb, title,
           ["Направление", "Кластер", "Артикул", "Отгрузки 14 дн", "Отгрузки 30 дн", "Коэф. сезона", "Скорость, шт/дн",
            "Доступно", "В пути", "Дней запаса", "Срок L, дн", "Потребность", "К отправке", "Коробов",
            "Короба", "Приоритет", "Пометки"],
           [[ln.warehouse, ln.cluster, ln.offer, ln.sales14, ln.sales30, ln.season, round(ln.speed, 2),
             ln.available, ln.transit, round(ln.days, 1) if ln.days is not None else None, ln.lead, ln.need,
             ln.qty, ln.box_count, _boxes(ln), ln.priority, "; ".join(ln.flags)] for ln in lines],
           [28, 24, 26, 9, 9, 8, 9, 9, 8, 9, 8, 10, 10, 8, 34, 10, 44])


def build_direction(plan, name, path, snap_time=""):
    """Файл одного направления: лист распределения и как посчитано — после «Всё верно» на экране направления."""
    wb = Workbook()
    wb.remove(wb.active)
    lines = [x for x in plan.lines if x.warehouse == name and x.qty > 0]
    _distribution(wb, plan, {name: lines})
    _need(wb, lines, "Как посчитано")
    _sheet(wb, "Параметры", ["Параметр", "Значение"], [
        ["Направление", name],
        ["Снимок данных", snap_time],
        ["Дата отгрузки", plan.ship_date.strftime("%d.%m.%Y")],
        ["Спрос", f"отгрузки FBO из кластера за {config.SALES_DAYS} дн без отмен ÷ {config.SALES_DAYS} × коэф. сезона"],
        ["Потребность", f"скорость × (срок L + {config.TARGET_DAYS}) − (доступно + в пути)"],
        ["К отправке", "потребность в пределах свободного остатка нашего склада, целыми моно-коробами"],
    ], [24, 90])
    wb.save(path)
    return path


def build(plan, path, snap_time=""):
    wb = Workbook()
    wb.remove(wb.active)

    _distribution(wb, plan, plan.by_warehouse())
    _need(wb, plan.lines)

    _sheet(wb, "Пустые кластеры",
           ["Кластер", "Артикул", "Продажи 14 дн (весь Ozon)", "Рекомендуемый склад", "Свободно у нас"],
           [[x["cluster"], "ни одного нашего товара в кластере", None, x["warehouse"], None] for x in plan.empty_whole]
           + [[x["cluster"], x["offer"], x["sales14"], x["warehouse"], x["ff_free"]] for x in plan.empty_clusters],
           [26, 28, 14, 30, 12])

    _sheet(wb, "Не попало в план", ["Артикул", "Направление", "Потребность", "Скорость, шт/дн", "Мин. короб", "Почему"],
           [[x["offer"], x["warehouse"], x["need"], round(x["speed"], 2), x["box"], x["why"]]
            for x in sorted(plan.small, key=lambda x: -x["need"])], [28, 28, 12, 12, 10, 26])

    _sheet(wb, "Решить - стоп", ["Артикул", "Статус", "Рейтинг", "Выкуп", "Продажи 14 дн", "Причина"],
           [[x["offer"], x.get("status"), x.get("rating"), x.get("buyout"), x.get("sales14"), x["reason"]]
            for x in plan.decide], [28, 16, 9, 9, 12, 70])

    _sheet(wb, "Заказать",
           ["Артикул", "Корень", "Заказать, шт", "На Ozon (доступно + в пути)", "Хватит на Ozon, дн", "Продажи, шт/дн",
            "Потребность кластеров", "Свободно у нас", "Отгружаем сейчас", "Прибыль, ₽/шт", "В заказ"],
           [[x["offer"], x["root"], x["order"], x["ozon"], round(x["days"], 1), round(x["speed"], 2), x["need"], x["free"],
             x["ship"], x["profit"], "да"] for x in plan.order]
           + [[x["offer"], x["root"], x["order"], x["ozon"], round(x["days"], 1), round(x["speed"], 2), x["need"], x["free"],
               x["ship"], x["profit"], "нет: " + x["why"]] for x in plan.order_skipped],
           [26, 20, 12, 14, 12, 12, 14, 12, 12, 12, 30])

    _sheet(wb, "Скоро платно", ["Артикул", "Шт", "Платное с", "Дней до платного"],
           [[x["offer"], x["qty"], x["date"], x["days"]] for x in plan.paid], [28, 8, 12, 14])

    _sheet(wb, "FBO+FBS", ["Артикул", "FBO, шт", "FBS, шт", "Сигнал"],
           [[x["offer"], x["fbo"], x["fbs"], "приоритет FBO — снять FBS"] for x in plan.fbo_fbs], [28, 10, 10, 28])

    _sheet(wb, "Баланс нашего склада",
           ["Артикул «Моего склада»", "Остаток − резерв", "В заявках, не уехали", "Свободно", "Отправляем", "Остаётся"],
           [[r["article"], r["stock"], r["unshipped"], r["free"], r["used"], r["free"] - r["used"]]
            for r in sorted(plan.ff.values(), key=lambda r: -r["used"]) if r["used"] or r["unshipped"]],
           [26, 14, 14, 12, 12, 12])

    _sheet(wb, "Параметры", ["Параметр", "Значение"], [
        ["Снимок данных", snap_time],
        ["Дата отгрузки", plan.ship_date.strftime("%d.%m.%Y")],
        ["Спрос", f"отгрузки FBO из кластера за {config.SALES_DAYS} дн без отмен ÷ {config.SALES_DAYS} × коэф. сезона"],
        ["Потребность", f"спрос × (L + {config.TARGET_DAYS}) − (доступно + в пути)"],
        ["Наш склад", f"«Мой склад», {config.MS_STORE}: остаток − резерв − заявки, которые ещё не уехали"],
        ["Короба", "моно-короб под одно направление; большими коробами, остаток — коробом поменьше"],
        ["Направления", "кластер целиком; склад приёмки — подсказка (наш основной склад кластера)"],
        ["Прибыль (Заказать)", "бот считает сам: с 1-го числа прошлого месяца по вчера. Продажи и комиссия — отчёт о "
                               "реализации Ozon, услуги Ozon — на штуку и по доле выручки, себестоимость — «Мой склад», "
                               "УСН 7%. В заказ — только прибыль на 1 шт больше 0"],
        ["Статусы", "таблица AlisaFox, лист «Юнит экономика Алиса», столбец «Статус ЗАЯВОК» (только чтение)"],
        ["Предупреждения", f"платное ≤ {config.PAID_SOON_DAYS} дн, выкуп < {config.BUYOUT_MIN:.0%}, "
                           f"рейтинг < {config.RATING_MIN}, сезон < {config.SEASON_LOW}"],
    ] + [["Сезон", f"{x['offer']} ({x['category']}): {x['k']} → {x['k_next']}"] for x in plan.season_notes]
      + [["Дубли", f"{x['root']}: {', '.join(x['offers'])}"] for x in plan.duplicates], [24, 90])

    wb.save(path)
    return path
