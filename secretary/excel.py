"""«Скачать отчёт»: Excel за неделю — категории и артикулы менеджера с изменением к прошлой неделе."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from bot import store
from secretary import report

GOOD, BAD = PatternFill("solid", fgColor="E3F4E8"), PatternFill("solid", fgColor="FBE4E4")


def _sheet(ws, head, rows, delta_cols=()):
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
        for i in delta_cols:
            v = r[i]
            if isinstance(v, (int, float)) and abs(v) >= report.FLAT:
                ws.cell(ws.max_row, i + 1).fill = GOOD if v > 0 else BAD
    for i, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(h) + 2)
    ws.freeze_panes = "B2"


def _d(cur, old):
    d = report.change(cur, old)
    return round(d) if d is not None else None


def build(db, title, skus, cur_p, prev_p, loss, sig):
    wb = Workbook()
    ws = wb.active
    ws.title = "Категории"
    rows = []
    for _, (name, part) in report.categories(db, skus).items():
        c, p = report.kpis(report.totals(db, part, *cur_p)), report.kpis(report.totals(db, part, *prev_p))
        if not c["orders"] and not p["orders"]:
            continue
        rows.append([name, c["orders"], _d(c["orders"], p["orders"]), round(c["revenue"]), _d(c["revenue"], p["revenue"]),
                     c["views"], _d(c["views"], p["views"]), round(c["ctr"], 1), round(c["buyout"]),
                     round(sum(loss[s]["profit"] for s in part if s in loss)),
                     sum(1 for s in part if s in loss and loss[s]["profit"] < 0),
                     sum(1 for s in part if s in sig and sig[s]["ship"])])
    _sheet(ws, ["Категория", "Заказы, шт", "Δ шт, %", "Заказы, ₽", "Δ ₽, %", "Показы", "Δ показов, %", "CTR, %",
                "Выкуп, %", "Чистая прибыль, ₽", "В минусе, арт.", "Пора отгрузить, арт."], rows, (2, 4, 6))
    ws = wb.create_sheet("Артикулы")
    rows = []
    for s in sorted(skus):
        c, p = report.kpis(report.totals(db, [s], *cur_p)), report.kpis(report.totals(db, [s], *prev_p))
        if not c["orders"] and not p["orders"]:
            continue
        r = store.find_sku(db, s)
        x = loss.get(s) or {}
        g = sig.get(s) or {}
        rows.append([r["offer_id"] if r else s, (r["name"] if r else "")[:60], r["type_name"] if r else "",
                     c["orders"], _d(c["orders"], p["orders"]), round(c["revenue"]), c["views"], _d(c["views"], p["views"]),
                     round(c["ctr"], 1), round(c["cart"], 1), round(c["position"]) if c["position"] else None,
                     round(c["buyout"]), round(x["ads"]) if x else None, round(x["profit"]) if x else None,
                     "да" if g.get("ship") else "", "да" if g.get("resort") else ""])
    _sheet(ws, ["Артикул", "Название", "Категория", "Заказы, шт", "Δ шт, %", "Заказы, ₽", "Показы", "Δ показов, %",
                "CTR, %", "В корзину, %", "Позиция", "Выкуп, %", "Реклама, ₽", "Чистая прибыль, ₽", "Пора отгрузить",
                "Пересорт"], rows, (4, 7))
    ws = wb.create_sheet("О расчёте")
    for line in (f"{title}", f"Неделя {cur_p[0]:%d.%m.%Y}–{cur_p[1]:%d.%m.%Y} к {prev_p[0]:%d.%m}–{prev_p[1]:%d.%m}",
                 "Источник: Ozon Seller API и Performance API, себестоимость — МойСклад / юнитка.",
                 "Прибыль = реализация − услуги Ozon (доля) − реклама артикула − себестоимость +3% − УСН 7%.",
                 "Выкуп = (доставлено − возвраты) / (доставлено + отмены). CTR = переходы в карточку / показы."):
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
